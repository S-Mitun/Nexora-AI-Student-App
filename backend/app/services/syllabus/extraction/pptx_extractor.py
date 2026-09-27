import os
from typing import List
import pptx
from app.services.syllabus.extraction.base import BaseDocumentExtractor, ExtractedDocument, ExtractedPage, ExtractedSection
from app.services.syllabus.extraction.normalizer import clean_academic_text
from app.core.logging import logger


class PptxExtractor(BaseDocumentExtractor):
    """
    Extracts structured syllabus content slide-by-slide from Microsoft PowerPoint (.pptx) files.
    Preserves slide numbers, shape titles, body text, tables, and speaker notes.
    """

    def extract(self, file_path: str) -> ExtractedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PPTX file does not exist at {file_path}")

        prs = pptx.Presentation(file_path)
        total_slides = len(prs.slides)
        pages: List[ExtractedPage] = []
        sections: List[ExtractedSection] = []
        full_text_parts: List[str] = []

        logger.info(f"[PptxExtractor] Starting extraction for {file_path} ({total_slides} slides)")

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_text_blocks: List[str] = []
            slide_title = f"Slide {slide_num}"

            # Extract shapes & text
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = clean_academic_text(paragraph.text)
                        if text:
                            slide_text_blocks.append(text)
                            if not slide_title or slide_title == f"Slide {slide_num}":
                                if len(text) < 100:
                                    slide_title = f"Slide {slide_num}: {text}"

                # Extract table data in shapes
                elif shape.has_table:
                    table_rows: List[str] = []
                    for row in shape.table.rows:
                        row_text = " | ".join(clean_academic_text(cell.text) for cell in row.cells if cell.text.strip())
                        if row_text:
                            table_rows.append(row_text)
                    if table_rows:
                        slide_text_blocks.append("\n[Slide Table]:\n" + "\n".join(table_rows))

            # Extract speaker notes if available
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = clean_academic_text(slide.notes_slide.notes_text_frame.text)
                    if notes:
                        slide_text_blocks.append(f"\n[Speaker Notes]: {notes}")
            except Exception as e:
                logger.debug(f"[PptxExtractor] Could not extract notes on slide {slide_num}: {e}")

            slide_text = "\n\n".join(slide_text_blocks)
            if slide_text:
                full_text_parts.append(f"--- Slide {slide_num} ---\n{slide_text}")

            pages.append(
                ExtractedPage(
                    page_number=slide_num,
                    raw_text=slide_text,
                    clean_text=slide_text,
                    has_images=any(s.shape_type == pptx.enum.shapes.MSO_SHAPE.PICTURE for s in slide.shapes if hasattr(s, "shape_type")),
                    is_scanned=False,
                    word_count=len(slide_text.split()),
                    metadata={"slide_number": slide_num},
                )
            )

            sections.append(
                ExtractedSection(
                    section_id=f"slide_{slide_num}",
                    section_title=slide_title,
                    page_start=slide_num,
                    page_end=slide_num,
                    raw_text=slide_text,
                    clean_text=slide_text,
                    structural_type="slide_block",
                    evidence_snippet=slide_text[:200] if slide_text else None,
                )
            )

        full_text = "\n\n".join(full_text_parts)

        return ExtractedDocument(
            source_path=file_path,
            file_type="pptx",
            total_pages=total_slides,
            pages=pages,
            sections=sections,
            full_text=full_text,
            has_scanned_pages=False,
            extraction_errors=[],
        )
